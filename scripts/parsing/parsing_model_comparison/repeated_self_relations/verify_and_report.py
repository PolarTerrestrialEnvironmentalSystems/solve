from pathlib import Path
import json
import sys
from collections import Counter
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT.parent/'baseline'))
from common import load_env,write_json,digest,file_hash
from schemas import Event,Relationship
from openpyxl import load_workbook

def read(p):return json.loads(p.read_text(encoding='utf-8'))
def lines(p):return [json.loads(s) for s in p.read_text(encoding='utf-8').splitlines() if s.strip()]

results=[read(p) for p in sorted((ROOT/'results').glob('*/result.json'))]
assert len(results)==18,'Study incomplete'
plan=read(ROOT/'study_plan.json')
summary=read(ROOT/'summary.json')
assert Counter((r['variant'],r['repeat'],r['case']) for r in results)==Counter((v,n,c['chunk_id']) for v in ['reference','compact'] for n in [1,2,3] for c in plan['cases'])
manifests=[read(ROOT/'results'/r['name']/'manifest.json') for r in results]
assert len({m['config']['schema_hash'] for m in manifests})==1
assert len({m['config']['ontology_hash'] for m in manifests})==1
assert len({digest(m['config']['settings']) for m in manifests})==1
assert len({digest(m['config']['model_identity']) for m in manifests})==1
assert len({m['config']['controller_hash'] for m in manifests})==1
for v in ['reference','compact']:
    assert len({m['config']['prompt_hash'] for m in manifests if m['config']['variant']==v})==1
    assert len({digest(m['config']['code_hashes']) for m in manifests if m['config']['variant']==v})==1
for m in manifests:
    c=m['config']; assert c['case']==plan['cases'][c['case_index']]
    assert c['plan_hash']==digest(plan)
    assert c['cover']==plan['cover']
reasons={}
for r in results:
    folder=ROOT/'results'/r['name']
    events=lines(folder/'fein/events.jsonl');rels=lines(folder/'fein/relationships.jsonl')
    ids={e['event_id'] for e in events}
    assert all(e['source_event_id'] in ids and e['target_event_id'] in ids for e in rels)
    assert sum(e['source_event_id']==e['target_event_id'] for e in rels)==r['exported_self_relationships']
    raw=[e for p in r['packages'] if p['status']=='complete' for e in p['result']['relationships']]
    assert sum(e['source_event_id']==e['target_event_id'] for e in raw)==r['raw_self_relationships']
    for name,model,rows in [('events',Event,events),('relationships',Relationship,rels)]:
        book=load_workbook(folder/'fein'/f'{name}.xlsx',read_only=True)
        assert list(next(book.active.values))==list(model.model_fields)
        assert book.active.max_row==len(rows)+1
        book.close()
    reasons[r['name']]=dict(Counter(reason for row in lines(folder/'fein/review_items.jsonl') for reason in row.get('reasons',[])))
env={};load_env(Path(r'C:\Users\jowals001\awi\solve\scripts\graphrag\data\.env'),env)
key=env['GRAPHRAG_API_KEY2'].encode('utf-8')
for p in ROOT.rglob('*'):
    if p.is_file() and p.suffix.lower() in {'.json','.jsonl','.md','.csv','.txt','.py'}:
        assert key not in p.read_bytes(),'Secret found; do not deliver'
verification={'all_18_runs_present':True,'identical_settings_schema_and_inputs':True,'unchanged_prompts_within_variants':True,
              'recomputed_self_edge_counts':True,'valid_export_references':True,'excel_jsonl_consistency':True,'secret_scan_passed':True}
write_json(ROOT/'final_verification.json',verification)
write_json(ROOT/'review_reason_counts.json',reasons)

a,b=summary['reference'],summary['compact']
def cell(s,key,total='runs_recorded'):return f"{s[key]}/{s[total]}"
report=f'''# Wiederholungstest: Selbstbeziehungen in Qwen-Antworten

Getestet am 28.09.2026 mit der zweiten Nutzer-PDF: Bericht von Manfred Simon und
Jürgen Böhme zu vereinbarten minimalen mittleren Monatsabflüssen der Elbe am
Grenzprofil Hřensko/Schöna, März 2012. Keine Prüfung der heutigen hydrologischen
oder rechtlichen Gültigkeit der historischen Aussagen.

## Beobachtetes Ergebnis

| Kriterium | Bisheriger Prompt | Kompakter Prompt |
|---|---:|---:|
| Aufgezeichnete Läufe | {a['runs_recorded']} | {b['runs_recorded']} |
| Vollständig schema-gültige Läufe | {a['successful_runs']} | {b['successful_runs']} |
| Läufe mit Selbstbeziehung in schema-gültiger Modellantwort | {cell(a,'runs_with_raw_self')} | {cell(b,'runs_with_raw_self')} |
| Selbstbeziehungen / alle Beziehungen vor Exportprüfung | {cell(a,'raw_self_edges','raw_edges')} | {cell(b,'raw_self_edges','raw_edges')} |
| Läufe mit Selbstbeziehung im finalen Export | {cell(a,'runs_with_exported_self')} | {cell(b,'runs_with_exported_self')} |
| Selbstbeziehungen / alle exportierten Beziehungen | {cell(a,'exported_self_edges','exported_edges')} | {cell(b,'exported_self_edges','exported_edges')} |
| Läufe ohne Beziehung in Modellantwort | {a['zero_raw_edge_runs']} | {b['zero_raw_edge_runs']} |
| Läufe ohne exportierte Beziehung | {a['zero_exported_edge_runs']} | {b['zero_exported_edge_runs']} |
| HTTP-Versuche | {a['http_attempts']} | {b['http_attempts']} |
| Ausgabetokens laut API | {a['completion_tokens']} | {b['completion_tokens']} |
| Summe Extraktionsdauer | {a['elapsed_seconds']:.1f} s | {b['elapsed_seconds']:.1f} s |

Ein Lauf zählt bereits bei einer Selbstbeziehung als betroffen. Die Kantenanzahl
ist zusätzlich angegeben, damit unterschiedlich viele erzeugte Beziehungen nicht
verborgen bleiben. „Schema-gültig“ bedeutet nicht fachlich richtig oder exportierbar.

## Aufteilung nach Textausschnitt

| Ausschnitt | Bisherig: betroffene Modellantworten | Kompakt: betroffene Modellantworten | Bisherig: betroffene Exporte | Kompakt: betroffene Exporte |
|---|---:|---:|---:|---:|
'''
labels={'moldau_qm364':'Moldau, Talsperrenwirkung (Seite 10)','eger_elbe':'Eger und Elbe, Niedrigwasseraufhöhung (Seite 11)','seasonal_climate':'Saisonaler Niederschlag und Abfluss (Seite 17)'}
for case in labels:
    x=a['cases'][case];y=b['cases'][case]
    report+=f"| {labels[case]} | {x['raw_self_runs']}/{x['runs']} | {y['raw_self_runs']}/{y['runs']} | {x['exported_self_runs']}/{x['runs']} | {y['exported_self_runs']}/{y['runs']} |\n"
report+='''
## Testbedingungen und Aussagegrenzen

Drei vorab ausgewählte kausale Passagen wurden je dreimal pro Variante verarbeitet
(18 Einzelläufe). Die Passagen stammen aus einer PDF, nicht aus 18 unabhängigen
Dokumenten. Textspannen, Modell, Schema, Ontologie und Einstellungen waren identisch.
Die Variantenreihenfolge wechselte; Antworten wurden nicht aus Checkpoints übernommen.
Keine Anti-Selbstbeziehungsregel wurde ergänzt und keine Exportprüfung gelockert.

Modellbasis `nvidia/Qwen3.8-Flash-Next-NVFP4`, Temperatur 0, Thinking aus,
16.000 Ausgabetokens, eine Anfrage gleichzeitig, gemeinsamer 30-RPM-Limiter.
Die genaue Modell-ID und alle Prüfsummen stehen in den Laufmanifesten.
Python 3.14.7 mit den gesperrten Paketversionen des ersten Tests.

Wiederholungen können bei Temperatur 0 gleich ausfallen. Die Anzahl unterschiedlicher
normalisierter Antworten pro Text und Variante steht in `summary.json`. Daher
stellen die Häufigkeiten nur Beobachtungen dieser gezielten Stichprobe dar; daraus
folgt keine allgemeine Fehlerwahrscheinlichkeit oder statistisch abgesicherte
Überlegenheit einer Promptvariante. Der erste Pilot wird nicht mit diesen Zahlen
zusammengerechnet, weil andere Dokumente und Ausschnitte verwendet wurden.

Die Modellstufe zählt Beziehungen aus erfolgreich validierten Extraktionspaketen.
Rohantworten vor etwaiger Reparatur oder abgebrochene Antworten bleiben separat
gespeichert. Die endgültige Exportstufe kann Beziehungen zusätzlich verlieren,
wenn Events, Themen-IDs oder Belege ungültig sind. Ein leerer Export belegt deshalb
keine erfolgreiche Behandlung von Selbstbeziehungen.

## Dateien

- `study_plan.json`: vorab festgelegte Textspannen und Auswertungseinheit.
- `summary.json` und `run_metrics.csv`: aggregierte Zahlen und alle Einzelläufe.
- `self_relation_examples.json`: betroffene Beziehungen und zugehörige Events.
- `review_reason_counts.json`: Gründe der bestehenden Qualitätsprüfungen.
- `results/`: Rohantworten, Checkpoints, normalisierte Pakete, Exporte, Prüflisten.
- `final_verification.json`: gleiche Eingaben, stabile Prompts, neu gezählte
  Selbstbeziehungen, Referenz- und Excel-Prüfung sowie Geheimnisscan bestätigt.

Das Original auf Y: sowie beide getesteten Prompts blieben unverändert.
'''
(ROOT/'Testbericht.md').write_text(report,encoding='utf-8')
print('Verified 18 runs and wrote report.')
