from pathlib import Path
import json
import sys
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'compact'))
from common import load_env, write_json
from schemas import Event, Relationship
from openpyxl import load_workbook

def read(p): return json.loads(p.read_text(encoding='utf-8'))
configs=[read(ROOT/'results'/v/'manifest.json')['config'] for v in ['reference','compact_qwen','compact_gptoss']]
for key in ['cases','cover','schema_hash','ontology_hash','dependencies','workers']:
    assert all(c[key]==configs[0][key] for c in configs)
checks={}
for v in ['reference','compact_qwen','compact_gptoss']:
    folder=ROOT/'results'/v
    assert len(list(folder.glob('case_*.json')))==3
    for name,model in [('events',Event),('relationships',Relationship)]:
        book=load_workbook(folder/'fein'/f'{name}.xlsx',read_only=True)
        rows=[json.loads(s) for s in (folder/'fein'/f'{name}.jsonl').read_text(encoding='utf-8').splitlines()]
        assert list(next(book.active.values))==list(model.model_fields)
        assert book.active.max_row==len(rows)+1
        book.close()
    checks[v]={'three_cases_recorded':True,'xlsx_jsonl_consistent':True}
secrets={}
load_env(Path(r'C:\Users\jowals001\awi\solve\scripts\graphrag\data\.env'),secrets)
key=secrets['GRAPHRAG_API_KEY2'].encode('utf-8')
for p in ROOT.rglob('*'):
    if not p.is_file() or '.venv' in p.parts or '__pycache__' in p.parts: continue
    if p.suffix.lower() in {'.json','.jsonl','.csv','.md','.txt','.py','.diff','.env'}:
        assert key not in p.read_bytes(), 'Secret found in generated artifact; do not distribute.'
write_json(ROOT/'final_verification.json',{'matched_inputs':True,'checks':checks,'secret_scan_passed':True})
print('Identical benchmark inputs, all 9 case records, Excel/JSONL consistency and secret scan: OK')
