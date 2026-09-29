"""Typed extraction schema, in precisely the order of the supplied templates."""
from __future__ import annotations

from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, JsonValue

Object = dict[str, JsonValue]
Assertion = Literal["reported", "observed", "suspected", "predicted", "hypothetical"]
Polarity = Literal["affirmed", "negated"]
TimeStatus = Literal["explicit", "resolved_from_context", "unknown_in_source", "not_stated", "no_specific_time"]


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)


class Evidence(StrictModel):
    quote: str = Field(min_length=1)
    block_id: str = Field(min_length=1)
    page: int = Field(ge=1, description="1-basierter PDF-Seitenindex, nicht gedruckte Seite")
    field_names: list[str] = Field(min_length=1, description="Alle durch diesen Beleg gestützten Ausgabefelder")
    printed_page: str | None = None
    table_id: str | None = None
    row: int | None = Field(default=None, ge=1)
    column: str | None = None


class TimeValue(StrictModel):
    start: str | None = None
    end: str | None = None
    recurrence: str | None = None


class Quantity(StrictModel):
    parameter: str
    value_original: str
    value_normalized: float | str | None = None
    unit_original: str | None = None
    unit_normalized: str | None = None
    comparison_operator: Literal["=", "<", ">", "<=", ">=", "range", "approximate"] | None = None
    normalization_rule: str | None = None


class Identifier(StrictModel):
    scheme: str | None = None
    value: str
    valid_time: str | None = None
    issuer: str | None = None


class Participant(StrictModel):
    name_original: str
    role: str
    class_id: str | None = None
    identifiers: list[Identifier] | None = None


class WaterbodyReference(StrictModel):
    name_original: str
    role: str
    identifiers: list[Identifier] | None = None


class Assignment(StrictModel):
    status: Literal["resolved", "probable", "ambiguous", "unresolved"]
    reason: str | None = None


class PublicationDate(StrictModel):
    original: str
    normalized: str | None = None
    precision: Literal["day", "month", "year", "approximate", "range"] | None = None


class Event(StrictModel):
    event_id: str = Field(min_length=1, description="Lokale ID, z.B. e1; finale ID vergibt Python")
    record_type: Literal["Ereignis", "Zustand", "Messung", "Prozess", "Maßnahme", "Eigenschaft"]
    waterbody_references: list[WaterbodyReference] | None = None
    waterbody_name_original: str | None = None
    waterbody_identifiers: list[Identifier] | None = None
    waterbody_assignment: Assignment | None = None
    statement: str = Field(min_length=1)
    ontology_concept: str | None = None
    participants: list[Participant] | None = None
    spatial_scope: Object | None = None
    system_scope: Object | None = None
    values: list[Quantity] | None = None
    measurement_context: Object | None = None
    time_original: str | None = None
    time_status: TimeStatus = "not_stated"
    time_normalized: TimeValue | None = None
    time_precision: str | None = None
    assertion_mode: Assertion
    polarity: Polarity
    implementation_status: Literal["proposed", "planned", "approved", "started", "completed", "cancelled", "unknown", "not_applicable"] = "not_applicable"
    categories: list[str] = Field(default_factory=list)
    category_assessment: list[Object] | None = None
    source_document: Object | None = None
    publication_date: PublicationDate | None = None
    retrieved_at: str | None = None
    evidence: list[Evidence] = Field(min_length=1)
    extraction_metadata: Object | None = None
    quality_assessment: Object | None = None


class Relationship(StrictModel):
    relationship_id: str = Field(min_length=1)
    source_event_id: str = Field(min_length=1)
    target_event_id: str = Field(min_length=1)
    relationship_type: str = Field(min_length=1)
    statement: str = Field(min_length=1)
    relationship_scope: Literal["specific_case", "general_relationship"]
    waterbody_references: list[WaterbodyReference] | None = None
    spatial_scope: Object | None = None
    assertion_mode: Assertion
    polarity: Polarity
    derivation_status: Literal["explicit_in_source", "reconstructed_from_passages"] = "explicit_in_source"
    evidence_basis: str | None = None
    mechanism: str | None = None
    conditions: list[Object] | None = None
    additional_factors: Object | None = None
    effect_size: Object | None = None
    valid_time: Object | None = None
    response_time_definition: str | None = None
    response_time_anchors: Object | None = None
    response_time_value: Object | None = None
    response_time_origin: Literal["stated_in_source", "calculated", "not_stated"] = "not_stated"
    response_time_evidence: list[Evidence] | None = None
    effect_duration: Object | None = None
    categories: list[str] = Field(default_factory=list)
    source_document: Object | None = None
    publication_date: PublicationDate | None = None
    retrieved_at: str | None = None
    evidence: list[Evidence] = Field(min_length=1)
    extraction_metadata: Object | None = None
    quality_assessment: Object | None = None


class DocumentInfo(StrictModel):
    title: str | None = None
    issuer: str | None = None
    authors: list[str] | None = None
    publication_date: PublicationDate | None = None
    evidence: list[Evidence] = Field(default_factory=list)


class ContextRequest(StrictModel):
    reason: str
    search_terms: list[str] = Field(min_length=1, max_length=5)
    event_ids: list[str] = Field(default_factory=list)


class Extraction(StrictModel):
    document: DocumentInfo | None = None
    events: list[Event]
    relationships: list[Relationship]
    context_requests: list[ContextRequest] = Field(default_factory=list, max_length=3)
    notes: list[str] = Field(default_factory=list)


def read_headers(path: Path, model: type[BaseModel]) -> dict[str, str]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [(str(a).strip(), str(b or "").strip()) for a, b in
                book.active.iter_rows(min_col=1, max_col=2, values_only=True) if a]
    finally:
        book.close()
    if [key for key, _ in rows] != list(model.model_fields):
        raise ValueError(f"Feldnamen/Reihenfolge in {path.name} stimmen nicht mit schemas.py überein.")
    return dict(rows)


def ontology_profile(ontology: dict) -> dict:
    """No fictional examples or crawling vocabulary in the extraction prompt."""
    return {
        "id": ontology["id"], "version": ontology["version"],
        "classes": ontology["classes"], "relations": ontology["relations"],
        "topics": [{k: v for k, v in t.items() if k not in {"search_terms", "source_hints"}}
                   for t in ontology["topics"]],
        "controlled_values": ontology["controlled_values"],
        "parsing_profile": ontology["parsing_profile"],
        "detail_field_definitions": ontology.get("detail_field_definitions", {}),
    }


def relationship_types(ontology: dict) -> set[str]:
    return set(ontology["controlled_values"]["relation_kind"]) | {"core:associated_with", "core:precedes"}
