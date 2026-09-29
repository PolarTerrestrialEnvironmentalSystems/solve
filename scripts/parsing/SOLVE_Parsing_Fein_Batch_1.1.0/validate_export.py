"""Evidence checks, reference resolution, conservative deduplication and exports."""
from __future__ import annotations

import csv
import io
import re
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from common import PARSER_VERSION, PROMPT_VERSION, SCHEMA_VERSION, atomic_text, canonical, digest, write_json, write_jsonl
from prepare_document import normalize_quote
from schemas import Event, Relationship, relationship_types
from relationship_guard import is_causal_self_relationship

TECHNICAL = {"event_id", "relationship_id", "source_event_id", "target_event_id", "source_document",
             "publication_date", "retrieved_at", "extraction_metadata", "quality_assessment", "evidence"}


def present(value) -> bool:
    return value is not None and value != "" and value != [] and value != {}


def check_evidence(evidence: list[dict], blocks: dict) -> list[str]:
    errors = []
    for item in evidence:
        block = blocks.get(item["block_id"])
        if block is None:
            errors.append("evidence_block_not_supplied:" + item["block_id"])
        elif item["page"] != block["page"]:
            errors.append("evidence_page_mismatch:" + item["block_id"])
        elif normalize_quote(item["quote"]) not in normalize_quote(block["text"]):
            errors.append("quote_not_found:" + item["block_id"])
    if not evidence:
        errors.append("missing_evidence")
    return sorted(set(errors))


def checks(row: dict, ontology: dict, blocks: dict, event: bool) -> tuple[list[str], list[str]]:
    fatal = check_evidence(row["evidence"], blocks)
    warnings = []
    class_ids = {c["id"] for c in ontology["classes"]}
    topic_ids = {t["id"] for t in ontology["topics"]}
    if event and row["ontology_concept"] and row["ontology_concept"] not in class_ids:
        fatal.append("unknown_ontology_class")
    if event and not row["ontology_concept"]:
        warnings.append("needs_ontology_extension")
    if event and row["ontology_concept"] in class_ids:
        parents = {c["id"]: c.get("parent_class") for c in ontology["classes"]}
        ancestors, current = set(), row["ontology_concept"]
        while current and current not in ancestors:
            ancestors.add(current)
            current = parents.get(current)
        expected = {"Ereignis": {"core:Event"}, "Maßnahme": {"core:Intervention"},
                    "Zustand": {"core:State", "core:Assessment"}, "Messung": {"core:Observation"},
                    "Prozess": {"core:Process"}, "Eigenschaft": {"core:Entity", "core:State"}}
        if not ancestors.intersection(expected[row["record_type"]]):
            warnings.append("record_type_ontology_mismatch")
    if any(t not in topic_ids for t in row["categories"]):
        fatal.append("unknown_topic")
    if event:
        for participant in row["participants"] or []:
            if participant["class_id"] and participant["class_id"] not in class_ids:
                fatal.append("unknown_participant_class")
    elif row["relationship_type"] not in relationship_types(ontology):
        fatal.append("unknown_relationship_type")
    field_evidence = {name for e in row["evidence"] for name in e["field_names"]}
    empty_statuses = {"not_stated", "not_applicable", "unresolved", "unknown"}
    for name, value in row.items():
        if name in TECHNICAL or not present(value) or (isinstance(value, str) and value in empty_statuses):
            continue
        if name not in field_evidence:
            warnings.append("field_evidence_missing:" + name)
    for item in row["evidence"]:
        for name in item["field_names"]:
            if name not in row:
                warnings.append("unknown_evidence_field:" + name)
    if event:
        time = row["time_normalized"]
        if time:
            for key in ["start", "end"]:
                value = time[key]
                if value and not valid_partial_date(value):
                    warnings.append("non_iso_or_invalid_time:" + key)
            if time["start"] and time["end"] and len(time["start"]) == len(time["end"]) and time["start"] > time["end"]:
                warnings.append("time_interval_reversed")
        if row["time_status"] == "explicit" and not row["time_original"]:
            warnings.append("explicit_time_without_original")
        quote_text = normalize_quote(" ".join(e["quote"] for e in row["evidence"]))
        for value in row["values"] or []:
            if normalize_quote(value["value_original"]) not in quote_text:
                warnings.append("quantity_original_not_in_evidence")
            if value["unit_normalized"] and value["unit_original"] != value["unit_normalized"] and not value["normalization_rule"]:
                warnings.append("unit_normalization_without_rule")
    else:
        if row["response_time_evidence"]:
            fatal.extend(check_evidence(row["response_time_evidence"], blocks))
        if row["response_time_value"] and not row["response_time_evidence"]:
            warnings.append("response_time_without_specific_evidence")
        if row["response_time_origin"] == "calculated" and not (row["response_time_anchors"] or {}).get("calculation"):
            warnings.append("response_time_calculation_missing")
    # Correct locator metadata comes from the parser, not the model.
    for item in row["evidence"] + (row.get("response_time_evidence") or []):
        if item["block_id"] in blocks:
            block = blocks[item["block_id"]]
            item["printed_page"] = block.get("printed_page")
            item["table_id"] = block.get("table_id")
            item["row"] = block.get("row")
    return sorted(set(fatal)), sorted(set(warnings))


def valid_partial_date(value: str) -> bool:
    if not re.fullmatch(r"\d{4}(?:-\d{2}(?:-\d{2})?)?", value):
        return False
    try:
        date.fromisoformat(value + {4: "-01-01", 7: "-01", 10: ""}[len(value)])
        return True
    except ValueError:
        return False


def document_metadata(packages: list[dict], document: dict, review: list[dict]) -> dict:
    result = {**document, "metadata_evidence": []}
    for package in packages:
        if package["status"] != "complete" or not package["result"].get("document"):
            continue
        doc = package["result"]["document"]
        by_id = {b["block_id"]: b for b in package["supplied_blocks"]}
        if check_evidence(doc["evidence"], by_id):
            review.append({"kind": "document_metadata", "reason": "invalid_evidence", "candidate": doc})
            continue
        supported = {k for e in doc["evidence"] for k in e["field_names"]}
        for key in ["title", "issuer", "authors", "publication_date"]:
            if not present(doc[key]):
                continue
            if key == "publication_date":
                quotes = normalize_quote(" ".join(e["quote"] for e in doc["evidence"] if key in e["field_names"]))
                if normalize_quote(doc[key]["original"]) not in quotes:
                    review.append({"kind": "document_metadata", "reason": "date_original_not_in_evidence", "candidate": doc[key]})
                    continue
                normalized = doc[key].get("normalized")
                if normalized and not valid_partial_date(normalized):
                    review.append({"kind": "document_metadata", "reason": "invalid_publication_date", "candidate": doc[key]})
                    continue
            if key not in supported:
                review.append({"kind": "document_metadata", "reason": "field_evidence_missing", "field": key})
            elif result.get(key) is None:
                result[key] = doc[key]
            elif result[key] != doc[key]:
                review.append({"kind": "document_metadata", "reason": "conflicting_metadata", "field": key, "candidate": doc[key]})
        result["metadata_evidence"] = unique(result["metadata_evidence"] + doc["evidence"])
    return result


def unique(items: list) -> list:
    return list({canonical(item): item for item in items}.values())


def semantic_key(row: dict) -> str:
    excluded = {"event_id", "relationship_id", "evidence", "extraction_metadata", "quality_assessment",
                "source_document", "publication_date", "retrieved_at", "category_assessment"}
    return digest({k: v for k, v in row.items() if k not in excluded})


def enrich(row: dict, package: dict, ontology: dict, document: dict, signature: str,
           variant: str, warnings: list[str]) -> None:
    model_assessment = row.get("quality_assessment")
    if model_assessment:
        warnings = sorted(set(warnings + ["model_reported_quality_notes"]))
    row["source_document"] = {k: v for k, v in document.items() if k not in {"publication_date", "retrieved_at", "pdf_metadata"}}
    row["publication_date"] = document.get("publication_date")
    row["retrieved_at"] = document.get("retrieved_at")
    row["extraction_metadata"] = {"run_signature": signature, "parser_version": PARSER_VERSION,
                                  "prompt_version": PROMPT_VERSION, "schema_version": SCHEMA_VERSION,
                                  "ontology_id": ontology["id"], "ontology_version": ontology["version"],
                                  "ontology_variant": variant, "ontology_hash": digest(ontology),
                                  "model": package.get("response_model"), "extracted_at": package["extracted_at"],
                                  "chunks": [package["chunk_id"]], "response_files": [package["raw_response_file"]]}
    row["quality_assessment"] = {"review_status": "needs_review" if warnings else "extracted",
                                 "issues": warnings, "quotes_verified": True,
                                 "semantic_review_performed": False,
                                 "needs_ontology_extension": "needs_ontology_extension" in warnings,
                                 "model_assessment_unverified": model_assessment}


def merge_row(target: dict, row: dict) -> None:
    target["evidence"] = unique(target["evidence"] + row["evidence"])
    for field in ["chunks", "response_files"]:
        target["extraction_metadata"][field] = unique(target["extraction_metadata"][field] + row["extraction_metadata"][field])
    target["quality_assessment"]["issues"] = sorted(set(target["quality_assessment"]["issues"] + row["quality_assessment"]["issues"]))
    if target["quality_assessment"]["issues"]:
        target["quality_assessment"]["review_status"] = "needs_review"


def resolve_ref(value: str, chunk_id: str, references: dict) -> str | None:
    # Prefer the local namespace even if a local ID happens to contain a colon.
    return references.get(chunk_id + ":" + value, references.get(value))


def finalize(packages: list[dict], ontology: dict, document: dict, signature: str,
             variant: str, extra_review: list[dict] | None = None) -> tuple[list, list, list, dict]:
    review = list(extra_review or [])
    for package in packages:
        review.extend(package.get("relationship_review", []))
    document = document_metadata(packages, document, review)
    events_by_key, relations_by_key, references = {}, {}, {}
    duplicate_events = duplicate_relations = 0
    for package in packages:
        if package["status"] != "complete":
            review.append({"kind": "failed_chunk", **package})
            continue
        blocks = {b["block_id"]: b for b in package["supplied_blocks"]}
        for source in package["result"]["events"]:
            row = Event.model_validate(source).model_dump(mode="json")
            fatal, warnings = checks(row, ontology, blocks, True)
            if fatal:
                review.append({"kind": "rejected_event", "chunk_id": package["chunk_id"], "reasons": fatal, "candidate": row})
                continue
            original_id = row["event_id"]
            enrich(row, package, ontology, document, signature, variant, warnings)
            key = semantic_key(row)
            row["event_id"] = "ev_" + variant + "_" + digest({"document_run": signature, "semantic": key})[:24]
            references[package["chunk_id"] + ":" + original_id] = row["event_id"]
            if key in events_by_key:
                merge_row(events_by_key[key], row)
                duplicate_events += 1
            else:
                events_by_key[key] = row
    for package in packages:
        if package["status"] != "complete":
            continue
        blocks = {b["block_id"]: b for b in package["supplied_blocks"]}
        for source in package["result"]["relationships"]:
            row = Relationship.model_validate(source).model_dump(mode="json")
            fatal, warnings = checks(row, ontology, blocks, False)
            if is_causal_self_relationship(row):
                fatal.append("causal_self_relationship")
            for key in ["source_event_id", "target_event_id"]:
                resolved = resolve_ref(row[key], package["chunk_id"], references)
                if not resolved:
                    fatal.append("unresolved_" + key)
                else:
                    row[key] = resolved
            # Re-check after namespace resolution and event deduplication: two
            # different input IDs can resolve to the same final event.
            if is_causal_self_relationship(row):
                fatal.append("causal_self_relationship")
            factors = row["additional_factors"] or {}
            for field in ["cause_event_ids", "effect_event_ids"]:
                if field in factors:
                    if not isinstance(factors[field], list) or not all(isinstance(v, str) for v in factors[field]):
                        fatal.append("invalid_" + field)
                        continue
                    resolved = [resolve_ref(v, package["chunk_id"], references) for v in factors[field]]
                    if any(v is None for v in resolved):
                        fatal.append("unresolved_" + field)
                    else:
                        factors[field] = sorted(set(resolved))
            if factors.get("combination") == "joint_contribution":
                causes, effects = factors.get("cause_event_ids", []), factors.get("effect_event_ids", [])
                if not isinstance(causes, list) or not isinstance(effects, list) or len(causes) < 2 or not effects:
                    fatal.append("incomplete_joint_cause_group")
                elif row["source_event_id"] not in causes or row["target_event_id"] not in effects:
                    fatal.append("edge_outside_joint_cause_group")
                else:
                    factors["group_id"] = "group_" + digest({"causes": causes, "effects": effects,
                                                             "scope": row["relationship_scope"], "polarity": row["polarity"]})[:20]
            if fatal:
                review.append({"kind": "rejected_relationship", "chunk_id": package["chunk_id"], "reasons": sorted(set(fatal)), "candidate": row})
                continue
            enrich(row, package, ontology, document, signature, variant, warnings)
            key = semantic_key(row)
            row["relationship_id"] = "rel_" + variant + "_" + digest({"document_run": signature, "semantic": key})[:24]
            if key in relations_by_key:
                merge_row(relations_by_key[key], row)
                duplicate_relations += 1
            else:
                relations_by_key[key] = row
    events, relationships = list(events_by_key.values()), list(relations_by_key.values())
    # Same source/subject/parameter but differing content is a review candidate, not a merge.
    possible = defaultdict(list)
    for event in events:
        k = digest({"type": event["record_type"], "waterbody": event["waterbody_references"],
                    "evidence": sorted({e["block_id"] for e in event["evidence"]}),
                    "parameters": [v["parameter"] for v in event["values"] or []],
                    "time": event["time_normalized"]})
        possible[k].append(event["event_id"])
        if event["quality_assessment"]["issues"]:
            review.append({"kind": "event_needs_review", "event_id": event["event_id"], "reasons": event["quality_assessment"]["issues"]})
    for candidates in possible.values():
        if len(candidates) > 1:
            review.append({"kind": "possible_duplicate_or_conflict", "event_ids": candidates})
    for row in relationships:
        if row["quality_assessment"]["issues"]:
            review.append({"kind": "relationship_needs_review", "relationship_id": row["relationship_id"], "reasons": row["quality_assessment"]["issues"]})
    summary = {"variant": variant, "events": len(events), "relationships": len(relationships),
               "causal_self_relationships_rejected": sum(
                   "causal_self_relationship" in r.get("reasons", []) for r in review),
               "causal_self_relationship_repairs": dict(Counter(
                   r["status"] for r in review if r["kind"] == "causal_self_relationship_repair")),
               "complete_chunks": sum(p["status"] == "complete" for p in packages),
               "failed_chunks": sum(p["status"] != "complete" for p in packages),
               "event_types": dict(Counter(e["record_type"] for e in events)),
               "topics": dict(Counter(t for e in events for t in e["categories"])),
               "merged_event_duplicates": duplicate_events, "merged_relationship_duplicates": duplicate_relations,
               "review_items": len(review), "review_kinds": dict(Counter(r["kind"] for r in review)),
               "event_fields_filled": {k: sum(present(e[k]) for e in events) for k in Event.model_fields},
               "relationship_fields_filled": {k: sum(present(e[k]) for e in relationships) for k in Relationship.model_fields},
               "empty_chunks": [p["chunk_id"] for p in packages if p["status"] == "complete" and not p["result"]["events"] and not p["result"]["relationships"]],
               "model_notes": [{"chunk_id": p["chunk_id"], "notes": p["result"]["notes"]} for p in packages if p["status"] == "complete" and p["result"]["notes"]]}
    return events, relationships, review, summary


def export_xlsx(path: Path, rows: list[dict], headers: list[str], review: list[dict]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Events" if "event_id" in headers else "Relationships"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234E70")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, row in enumerate(rows, 2):
        for col, key in enumerate(headers, 1):
            value = row[key]
            if isinstance(value, (dict, list)):
                value = canonical(value)
            if isinstance(value, str) and (len(value.encode("utf-16-le")) // 2 > 32767 or ILLEGAL_CHARACTERS_RE.search(value)):
                value = f"Vollständiger Wert: {path.stem}.jsonl, Datensatz {index-1}, Feld {key}"
                review.append({"kind": "excel_value_externalized", "file": path.name, "row": index, "field": key})
            cell = sheet.cell(index, col, value)
            if isinstance(value, str):
                # Literal source text must never become an Excel formula.
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        sheet.column_dimensions[cell.column_letter].width = 28 if cell.value not in {"statement", "evidence"} else 65
    sheet.row_dimensions[1].height = 32
    for index in range(2, len(rows)+2):
        sheet.row_dimensions[index].height = 60
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(".tmp.xlsx")
    book.save(temp)
    book.close()
    temp.replace(path)
    # Reopen to check the externally visible schema and data row count.
    check = load_workbook(path, read_only=True, data_only=False)
    try:
        if list(next(check.active.iter_rows(min_row=1, max_row=1, values_only=True))) != headers:
            raise ValueError("Excel-Headerprüfung fehlgeschlagen.")
        if check.active.max_row != len(rows)+1:
            raise ValueError("Excel-Zeilenzahl stimmt nicht mit JSONL überein.")
    finally:
        check.close()


def export_variant(output: Path, events: list, relationships: list, review: list, summary: dict, headers: dict) -> None:
    write_jsonl(output / "events.jsonl", events)
    write_jsonl(output / "relationships.jsonl", relationships)
    export_xlsx(output / "events.xlsx", events, list(headers["events"]), review)
    export_xlsx(output / "relationships.xlsx", relationships, list(headers["relationships"]), review)
    summary["review_items"] = len(review)
    summary["review_kinds"] = dict(Counter(r["kind"] for r in review))
    write_jsonl(output / "review_items.jsonl", review)
    write_json(output / "summary.json", summary)


def compare(output: Path, results: dict, fine_ontology: dict, coverage: dict, api_stats: dict) -> None:
    mapping = fine_ontology.get("coarse_mapping", {})
    classes = {x["fine_id"]: x["coarse_id"] for x in mapping.get("classes", [])}
    topics = {x["fine_id"]: x["coarse_id"] for x in mapping.get("topics", [])}
    def key(row):
        return digest({"type": row["record_type"], "waterbody": row["waterbody_references"],
                       "time": row["time_normalized"], "values": row["values"],
                       "evidence": sorted({(e["page"], e["block_id"], normalize_quote(e["quote"])) for e in row["evidence"]}),
                       "coarse_class": classes.get(row["ontology_concept"], row["ontology_concept"])})
    fine_index = defaultdict(list)
    for row in results.get("fein", {}).get("events", []):
        fine_index[key(row)].append(row)
    matches, used = [], set()
    for row in results.get("grob", {}).get("events", []):
        candidates = fine_index.get(key(row), [])
        if not candidates:
            matches.append({"grob_event_id": row["event_id"], "fein_event_id": "", "match": "no_exact_evidence_candidate"})
        for candidate in candidates:
            used.add(candidate["event_id"])
            matches.append({"grob_event_id": row["event_id"], "fein_event_id": candidate["event_id"],
                            "match": "evidence_candidate" if len(candidates) == 1 else "ambiguous_evidence_candidate"})
    for row in results.get("fein", {}).get("events", []):
        if row["event_id"] not in used:
            matches.append({"grob_event_id": "", "fein_event_id": row["event_id"], "match": "no_exact_evidence_candidate"})
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=["grob_event_id", "fein_event_id", "match"])
    writer.writeheader()
    writer.writerows(matches)
    atomic_text(output / "comparison.csv", stream.getvalue())
    text = ["# Vergleich der Parsing-Läufe", "", "Automatische Kennzahlen; keine manuell bestimmte Precision/Recall.", "",
            f"PDF-Seiten: {coverage['total_pages']}; ausgewählt: {len(coverage['selected_pages'])}; Abschnitte: {coverage['chunks']}.",
            "Bildinhalte wurden nicht interpretiert. Details und Tabellenhinweise: prepared/coverage.json.", "",
            "| Lauf | Events | Beziehungen | Fehlgeschlagene Abschnitte | Prüffälle |", "| --- | ---: | ---: | ---: | ---: |"]
    details = {}
    for name, result in results.items():
        s = result["summary"]
        text.append(f"| {name} | {s['events']} | {s['relationships']} | {s['failed_chunks']} | {s['review_items']} |")
        details[name] = {"summary": s, "coarse_topics": dict(Counter(t for e in result["events"] for t in {topics.get(c, c) for c in e["categories"]}))}
    text += ["", f"API-Versuche insgesamt: {api_stats.get('requests', 0)}; maximales gleitendes 60-s-Fenster: {api_stats.get('max_requests_last_60s', 0)}.",
             "Minutenstatistik: logs/api_requests_per_minute.csv. Rohdaten und Laufzeiten: logs/api_requests.jsonl.",
             "", "comparison.csv enthält konservative Kandidaten anhand identischer Belege und Sachmerkmale.",
             "Kein Kandidat bedeutet nicht automatisch einen zusätzlichen oder fehlenden Fakt. Unterschiedliche Formulierungen müssen fachlich geprüft werden.",
             "Feldbelegungszahlen stehen in comparison.json; optionale Felder sind nicht bei allen Datensätzen anwendbar.",
             "Ein vollständiger technischer Durchlauf ist kein Nachweis vollständiger oder fachlich korrekter Extraktion."]
    write_json(output / "comparison.json", {"variants": details, "api": api_stats, "coverage": coverage})
    atomic_text(output / "comparison.md", "\n".join(text) + "\n")
